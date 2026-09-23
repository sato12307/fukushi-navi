"""muni-master.py — 全市区町村（1,741）の一覧に生活保護の級地を付けて data/muni-master.json に書く。

  python scripts/muni-master.py         → data/muni-master.json
  python scripts/muni-master.py --dry   → 書かずに件数と突き合わせだけ出す

★元は2つだけ
  ① 総務省「全国地方公共団体コード」（.cache/hogo-shinsei/muni-code.xlsx・R6.1.1現在）
     1枚目のシート＝都道府県と市区町村（政令市は市の単位。区は2枚目にしか無い）。
  ② data/kyuchi.json（kyuchi-parse.py が厚労省PDFから作ったもの）。
     キーは「都道府県|市町村」の公式名。ここに無い市町村はすべて3級地-2。
     東京都の特別区は「東京都|区の存する地域」の1行で23区ぶん。

★止める条件
  - kyuchi.json のキーが1つでもコード表の市区町村に当たらない（名前のずれ＝級地の取りこぼし）
  - 市区町村の数が1,741でない
"""
import json
import os
import sys
import unicodedata

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MUNI = os.path.join(ROOT, ".cache", "hogo-shinsei", "muni-code.xlsx")
KYUCHI = os.path.join(ROOT, "data", "kyuchi.json")
OUT = os.path.join(ROOT, "data", "muni-master.json")
DRY = "--dry" in sys.argv


def die(msg):
    print(msg, file=sys.stderr)
    print("何も書き出していません。", file=sys.stderr)
    sys.exit(1)


wb = openpyxl.load_workbook(MUNI, read_only=True, data_only=True)
ws = wb[wb.sheetnames[0]]
rows = []
for r in ws.iter_rows(min_row=2, values_only=True):
    code, pref, city, _, kana = (list(r) + [None] * 5)[:5]
    if not code or not pref:
        continue
    code = str(code).strip()
    if not city:  # 都道府県の行
        continue
    # 北方領土の6村（色丹村・泊村・留夜別村・留別村・紗那村・蘂取村＝01695〜01700）はコード表に載るが、
    # 住民税を課す自治体として機能していない。古宇郡の泊村（01403）は別の実在の村なので名前でなくコードで外す。
    if "01695" <= code[:5] <= "01700":
        continue
    rows.append({
        "code": code[:5],          # 5桁（検査数字なし）＝URLに使う
        "code6": code,             # 6桁（検査数字つき）
        "pref": str(pref).strip(),
        "city": str(city).strip(),
        "kana": unicodedata.normalize("NFKC", str(kana or "")).strip(),
    })

if len(rows) != 1741:
    die(f"市区町村の数が1,741ではありません（{len(rows)}件）。コード表の版を確かめてください。")

ky = json.load(open(KYUCHI, encoding="utf-8"))["rows"]
by_key = {f'{r["pref"]}|{r["city"]}': r for r in rows}
used = set()
for key, grade in ky.items():
    if key == "東京都|区の存する地域":
        hits = [r for r in rows if r["pref"] == "東京都" and r["code"].startswith("131")]
        if len(hits) != 23:
            die(f"東京都の特別区が23件ではありません（{len(hits)}件）")
        for r in hits:
            r["kyuchi"] = grade
            used.add(f'{r["pref"]}|{r["city"]}')
        continue
    r = by_key.get(key)
    if not r:
        die(f"級地の表のキーがコード表に見つかりません: {key}")
    r["kyuchi"] = grade
    used.add(key)

for r in rows:
    r.setdefault("kyuchi", "3級地-2")
    # 住民税の非課税の線に効くのは「1級地・2級地・3級地」の3段だけ（-1・-2の区別は効かない）
    r["zeiKyuchi"] = r["kyuchi"][0]

from collections import Counter
cnt = Counter(r["kyuchi"] for r in rows)
print("級地の内訳:", dict(sorted(cnt.items())))
print("明示の級地:", sum(v for k, v in cnt.items() if k != "3級地-2"), "件 ／ 3級地-2（一覧に無い市町村）:", cnt["3級地-2"], "件")

if DRY:
    sys.exit(0)

out = {
    "note": "全市区町村の一覧（総務省 全国地方公共団体コード R6.1.1現在）に、生活保護の級地（厚生労働省 kyuchi.3010.pdf・平成30年10月1日現在）を付けたもの。zeiKyuchi は住民税の非課税の線に効く級地（1〜3）。scripts/muni-master.py が作る。手で直さない。",
    "sources": {
        "codes": "総務省 全国地方公共団体コード https://www.soumu.go.jp/denshijiti/code.html",
        "kyuchi": "厚生労働省 お住まいの地域の級地を確認 https://www.mhlw.go.jp/content/kyuchi.3010.pdf",
    },
    "rows": rows,
}
with open(OUT, "w", encoding="utf-8") as f:
    json.dump(out, f, ensure_ascii=False, indent=0)
print("書き出し:", OUT, len(rows), "件")
