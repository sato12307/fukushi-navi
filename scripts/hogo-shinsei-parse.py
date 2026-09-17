#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""被保護者調査 第18表・第9表を読んで、自治体別の申請/取下げ/却下の表にする。

  python scripts/hogo-shinsei-parse.py          # data/hogo-shinsei.json を作る
  python scripts/hogo-shinsei-parse.py --dry    # 書かずに結果だけ出す

前に回すもの: node scripts/hogo-shinsei-fetch.mjs

★この表の読み方（間違えると数字の意味が変わる）
  ・「都道府県」の行に **指定都市・中核市は入っていない**（表の注2）。
    つまり神奈川県の行は「横浜市・川崎市・相模原市・横須賀市…を除いた神奈川県」。
    県まるごとの数字ではないので、ページでもそう書く。
  ・市部／郡部の2列に割れている。合算して使う。
  ・年度累計（月分報告の積み上げ）。未処理件数だけは **年度末時点の残** で性質が違う。
  ・申請と開始は**別の月にまたがる**ので、開始÷申請はぴったり1にならない。
    さらに開始には申請によらない職権保護が混ざりうる。だから当サイトでは
    「開始率」を主役にせず、**却下率＝却下件数÷申請件数**を主役にする。

★年度はIDから決めない（e-Statの一覧は見出しと tstat の並びがずれる）。
  A1セルの「令和◯年度被保護者調査」を読んで、ファイル名の年度と一致するか検算する。
  一致しなければ止める。ここを通さないと1年ずれた表を平気で公開してしまう。

★人口は第9表から逆算する。
  第9表の注1＝「保護率＝1か月平均の被保護実人員数 ÷ 人口推計」。
  よって 人口 = 1か月平均 ÷ 保護率 × 1000。
  保護率は小数第1位までしか公表されないので、人口は**0.5%ほどの丸め誤差**を持つ。
  人口千人あたり申請件数は「多い/少ない」を見るためのもので、1位争いには使わない。
"""
import json
import os
import re
import sys
from glob import glob

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CACHE = os.path.join(ROOT, ".cache", "hogo-shinsei")
OUT = os.path.join(ROOT, "data", "hogo-shinsei.json")
DRY = "--dry" in sys.argv

# 第18表の列（0始まり）。市部・郡部の2列で1項目。
COLS18 = {
    "shinsei": 2,      # 申請件数
    "torisage": 4,     # 申請取下げ件数
    "kyakka": 6,       # 却下件数
    "kaishi_setai": 8,   # 保護開始世帯数
    "kaishi_nin": 10,    # 保護開始人員数
    "mishori": 12,     # 未処理件数（年度末現在）
    "haishi_setai": 14,  # 保護廃止世帯数
    "haishi_nin": 16,    # 保護廃止人員数
}
LEVEL = {"全国": "全国", "都道府県": "都道府県", "指定都市（別掲）": "指定都市", "中核市（別掲）": "中核市"}
WAREKI = re.compile(r"(令和|平成)([元0-9０-９]+)年度|^(\d{4})年度")
Z2H = str.maketrans("０１２３４５６７８９", "0123456789")


def nendo_from_a1(text):
    """A1の「令和６年度被保護者調査」→ 西暦の年度。読めなければ None。"""
    if not text:
        return None
    t = str(text).translate(Z2H).replace("六", "6")
    m = re.search(r"令和(元|\d+)年度", t)
    if m:
        n = 1 if m.group(1) == "元" else int(m.group(1))
        return n + 2018
    m = re.search(r"(\d{4})年度", t)
    if m:
        return int(m.group(1))
    # 「令和６年度」のような漢数字は openpyxl では全角数字で来る想定。来なければ諦める。
    return None


KANSUJI = {"元": 1, "一": 1, "二": 2, "三": 3, "四": 4, "五": 5, "六": 6, "七": 7, "八": 8, "九": 9, "十": 10}


def nendo_kansuji(text):
    m = re.search(r"令和([元一二三四五六七八九十])年度", str(text or ""))
    return KANSUJI[m.group(1)] + 2018 if m else None


def num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return int(v)
    s = str(v).replace(",", "").strip()
    # 「・」＝該当なし（指定都市・中核市の行には郡部が無い。島根県のように郡部の
    # 申請が0件の県にも入る）。「-」も0件。どちらも0として足してよいことは、
    # 全国行＝内訳の合計 が6年度ぶん一致することで確かめている（main の検算）。
    if s in ("-", "‐", "−", "・", "…", "･･･"):
        return None
    try:
        return int(float(s))
    except ValueError:
        return None


def read_rows(path, ncols):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = [r for r in ws.iter_rows(max_col=ncols, values_only=True)]
    a1 = rows[0][0] if rows else None
    return a1, rows


def parse_t18(path, want_year):
    a1, rows = read_rows(path, 22)
    got = nendo_from_a1(a1) or nendo_kansuji(a1)
    if got != want_year:
        raise SystemExit(f"年度がずれている: {os.path.basename(path)} は A1='{a1}' → {got} 年度。IDの対応表を直すこと。")
    out = {}
    for r in rows:
        lv = LEVEL.get(str(r[0]).strip() if r[0] else "")
        if not lv:
            continue
        name = str(r[1]).strip()
        rec = {"level": lv}
        for key, c in COLS18.items():
            a, b = num(r[c]), num(r[c + 1])
            rec[key] = (a or 0) + (b or 0)
        out[name] = rec
    return out


def parse_t9(path, want_year):
    """1か月平均の被保護実人員数と年度保護率（人口千対）→ 人口を逆算。"""
    a1, rows = read_rows(path, 18)
    got = nendo_from_a1(a1) or nendo_kansuji(a1)
    if got != want_year:
        raise SystemExit(f"年度がずれている: {os.path.basename(path)} は A1='{a1}' → {got} 年度。")
    out = {}
    for r in rows:
        lv = LEVEL.get(str(r[0]).strip() if r[0] else "")
        if not lv:
            continue
        name = str(r[1]).strip()
        avg = num(r[15])
        rate = r[16]
        try:
            rate = float(str(rate).replace(",", ""))
        except (TypeError, ValueError):
            rate = None
        pop = int(round(avg / rate * 1000)) if (avg and rate) else None
        out[name] = {"hogo_nin": avg, "hogo_ritsu": rate, "jinko": pop}
    return out


def city_to_pref():
    """総務省「全国地方公共団体コード」から 市名→都道府県名。

    ★同名の市があると対応が壊れる（例：府中市＝東京都と広島県）。
      だから重複した名前は **捨てる**。捨てた名前が第18表に出てきたら
      その場で止める（黙って別の県の名前を付けるより止めたほうがよい）。
    """
    path = os.path.join(CACHE, "muni-code.xlsx")
    if not os.path.exists(path):
        return {}, set()
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    m, dup = {}, set()
    for r in ws.iter_rows(min_row=2, max_col=3, values_only=True):
        pref, city = (str(r[1]).strip() if r[1] else ""), (str(r[2]).strip() if r[2] else "")
        if not city or not pref:
            continue
        if city in m and m[city] != pref:
            dup.add(city)
        m[city] = pref
    for c in dup:
        m.pop(c, None)
    return m, dup


def main():
    years = sorted(int(re.search(r"t18-(\d{4})", p).group(1)) for p in glob(os.path.join(CACHE, "t18-*.xlsx")))
    if not years:
        raise SystemExit("`.cache/hogo-shinsei/` が空。先に node scripts/hogo-shinsei-fetch.mjs を回すこと。")
    by_year = {y: parse_t18(os.path.join(CACHE, f"t18-{y}.xlsx"), y) for y in years}
    latest = years[-1]

    t9path = os.path.join(CACHE, f"t9-{latest}.xlsx")
    pop = parse_t9(t9path, latest) if os.path.exists(t9path) else {}

    # 件数の検算：全国行 = 都道府県47の合計 + 指定都市20 + 中核市62（別掲は都道府県に含まれない）
    for y in years:
        d = by_year[y]
        total = d.get("全国", {}).get("shinsei")
        parts = sum(v["shinsei"] for k, v in d.items() if v["level"] != "全国")
        if total != parts:
            raise SystemExit(f"{y}年度: 全国{total}件 ≠ 内訳合計{parts}件。表の読み方が違う。")
        n = {"都道府県": 0, "指定都市": 0, "中核市": 0}
        for v in d.values():
            if v["level"] in n:
                n[v["level"]] += 1
        if n["都道府県"] != 47:
            raise SystemExit(f"{y}年度: 都道府県が{n['都道府県']}件しかない。")
        print(f"  {y}年度 申請{total:,}件 / 都道府県{n['都道府県']} 指定都市{n['指定都市']} 中核市{n['中核市']}")

    c2p, dup = city_to_pref()
    kuridashi = {}   # 都道府県名 → その県から別掲されている市のリスト
    for name, rec in by_year[latest].items():
        if rec["level"] in ("指定都市", "中核市"):
            if name in dup:
                raise SystemExit(f"{name} は同名の市が複数あって都道府県を決められない。手で対応を足すこと。")
            pref = c2p.get(name)
            if not pref:
                raise SystemExit(f"{name} の都道府県が全国地方公共団体コードに無い。コード表が古い可能性。")
            kuridashi.setdefault(pref, []).append(name)

    rows = []
    for name, rec in by_year[latest].items():
        if name == "全国":
            continue
        r = dict(rec)
        r["name"] = name
        r["hist"] = {
            str(y): {
                "shinsei": by_year[y][name]["shinsei"],
                "kyakka": by_year[y][name]["kyakka"],
                "torisage": by_year[y][name]["torisage"],
            }
            for y in years if name in by_year[y]
        }
        r.update(pop.get(name, {}))
        if rec["level"] == "都道府県":
            r["excluded"] = sorted(kuridashi.get(name, []))
        else:
            r["pref"] = c2p.get(name)
        rows.append(r)

    national = dict(by_year[latest]["全国"])
    national["hist"] = {
        str(y): {k: by_year[y]["全国"][k] for k in ("shinsei", "kyakka", "torisage", "kaishi_setai")}
        for y in years
    }
    national.update(pop.get("全国", {}))

    data = {
        "source": "厚生労働省「被保護者調査（月次調査・確定値）」第18表・第9表",
        "latest_year": latest,
        "years": years,
        "national": national,
        "rows": sorted(rows, key=lambda r: r["name"]),
    }
    if DRY:
        k = national["kyakka"] / national["shinsei"] * 100
        t = national["torisage"] / national["shinsei"] * 100
        print(f"\n全国 {latest}年度: 申請{national['shinsei']:,} 却下{national['kyakka']:,}({k:.2f}%) 取下げ{national['torisage']:,}({t:.2f}%)")
        print(f"自治体 {len(rows)} 機関 / 人口逆算できた {sum(1 for r in rows if r.get('jinko'))} 機関")
        return
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=1)
    print(f"\n書いた: {OUT}  ({len(rows)}機関 × {len(years)}年度)")


if __name__ == "__main__":
    main()
