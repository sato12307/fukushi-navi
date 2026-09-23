# -*- coding: utf-8 -*-
# ─────────────────────────────────────────────────────────────────────────────
# hasan-parse.py — 司法統計 第4表と弁護士会別会員数を読んで data/hasan.json を書く。
#
#   python scripts/hasan-parse.py
#   前に回すもの: node scripts/hasan-fetch.mjs
#
# ★第4表の読み方（ここで一度やらかした）
#   列の見出し（裁判所名）は**3列ぶんの結合セル**で、その中身は 新受／既済／未済。
#   「その裁判所名の列を全部足す」と **新受＋既済＋未済** になり、破産が92,059件のところ
#   213,895件になる。名前が出ている先頭列（＝新受）だけを採る。
#
# ★地裁と都道府県の対応
#   50地裁のうち北海道だけ4つ（札幌・函館・旭川・釧路）。ほかは1県1地裁。
#   支部は第4表には出てこない（本庁の数字に含まれている）。
#
# ★関所
#   A. 各項目について **50地裁の合計＝同じ表の「全国総数」列** であること。1件でも違えば列がずれている。
#   B. 地裁の列がちょうど50本あること。
#   C. 弁護士会が52会そろい、合計が PDF の「合計」行と一致すること。
#   D. 都道府県が47そろうこと。
#
# ★年は暦年（1〜12月）。生活保護・困窮者支援の「年度」と同じ軸に並べない。
# ★人口と保護率は data/hogo-shinsei.json から都道府県ごとに足して作る
#   （保護率からの逆算なので0.5%ほどの丸め誤差を含む。順位争いには使わない）。
# ─────────────────────────────────────────────────────────────────────────────
import glob
import io
import json
import os
import re
import sys

import openpyxl
import pypdf

# Windows の既定のコンソール（cp932）は「↔」「—」「≤」などを encode できず、
# 進捗表示の1行で UnicodeEncodeError を出して途中で落ちる（2026-09-23 の川崎で実際に起きた）。
# 記号を選び直すのは漏れるので、出力の文字コードのほうを UTF-8 に固定する。
for _s in (sys.stdout, sys.stderr):
    if hasattr(_s, "reconfigure"):
        _s.reconfigure(encoding="utf-8", errors="replace")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CACHE = os.path.join(ROOT, '.cache', 'hasan')
HOGO = os.path.join(ROOT, 'data', 'hogo-shinsei.json')
OUT = os.path.join(ROOT, 'data', 'hasan.json')

# 地方裁判所 -> 都道府県（北海道だけ4地裁）
COURT2PREF = {
    '東京': '東京都', '横浜': '神奈川県', 'さいたま': '埼玉県', '千葉': '千葉県', '水戸': '茨城県',
    '宇都宮': '栃木県', '前橋': '群馬県', '静岡': '静岡県', '甲府': '山梨県', '長野': '長野県',
    '新潟': '新潟県', '大阪': '大阪府', '京都': '京都府', '神戸': '兵庫県', '奈良': '奈良県',
    '大津': '滋賀県', '和歌山': '和歌山県', '名古屋': '愛知県', '津': '三重県', '岐阜': '岐阜県',
    '福井': '福井県', '金沢': '石川県', '富山': '富山県', '広島': '広島県', '山口': '山口県',
    '岡山': '岡山県', '鳥取': '鳥取県', '松江': '島根県', '福岡': '福岡県', '佐賀': '佐賀県',
    '長崎': '長崎県', '大分': '大分県', '熊本': '熊本県', '鹿児島': '鹿児島県', '宮崎': '宮崎県',
    '那覇': '沖縄県', '仙台': '宮城県', '福島': '福島県', '山形': '山形県', '盛岡': '岩手県',
    '秋田': '秋田県', '青森': '青森県', '札幌': '北海道', '函館': '北海道', '旭川': '北海道',
    '釧路': '北海道', '高松': '香川県', '徳島': '徳島県', '高知': '高知県', '松山': '愛媛県',
}
# 弁護士会 -> 都道府県（東京3会・北海道4会・仙台=宮城・金沢=石川）
BAR2PREF = {
    '東京': '東京都', '第一東京': '東京都', '第二東京': '東京都', '神奈川県': '神奈川県',
    '埼玉': '埼玉県', '千葉県': '千葉県', '茨城県': '茨城県', '栃木県': '栃木県', '群馬': '群馬県',
    '静岡県': '静岡県', '山梨県': '山梨県', '長野県': '長野県', '新潟県': '新潟県', '大阪': '大阪府',
    '京都': '京都府', '兵庫県': '兵庫県', '奈良': '奈良県', '滋賀': '滋賀県', '和歌山': '和歌山県',
    '愛知県': '愛知県', '三重': '三重県', '岐阜県': '岐阜県', '福井': '福井県', '金沢': '石川県',
    '富山県': '富山県', '広島': '広島県', '山口県': '山口県', '岡山': '岡山県', '鳥取県': '鳥取県',
    '島根県': '島根県', '福岡県': '福岡県', '佐賀県': '佐賀県', '長崎県': '長崎県', '大分県': '大分県',
    '熊本県': '熊本県', '鹿児島県': '鹿児島県', '宮崎県': '宮崎県', '沖縄': '沖縄県', '仙台': '宮城県',
    '福島県': '福島県', '山形県': '山形県', '岩手': '岩手県', '秋田': '秋田県', '青森県': '青森県',
    '札幌': '北海道', '函館': '北海道', '旭川': '北海道', '釧路': '北海道', '香川県': '香川県',
    '徳島': '徳島県', '高知': '高知県', '愛媛': '愛媛県',
}
# 採る行（ラベル, 第3列の小分類, 出力キー）
WANT = [
    ('破産', None, 'hasan'),
    ('小規模個人再生', None, 'saisei'),
    ('給与所得者等再生', None, 'kyuyo'),
]
# ★「担保権の実行（不動産）」も同じ表から採れるが、記事では使わないので採らない。
#   持ち家仮説の傍証にならないか調べたが、再生を選ぶ割合との相関は r=+0.14 しかなく、
#   使わない列のために年を1つ落とす（2021年版はこの行のセルが壊れている）のは割に合わない。


class CellError(Exception):
    pass


def cellint(ws, r, c):
    """数字のセルを整数で読む。
    ★年によっては 2つの数字が改行で1セルに入っている行がある（取り込みミス）。
      int() が落ちるので、ここで捕まえてその年を丸ごと採らない。
      黙って0にすると件数では気づけないので、必ず例外にして止める。"""
    v = ws.cell(r, c).value
    if v is None or v == '':
        return 0
    if isinstance(v, (int, float)):
        return int(v)
    t = str(v).strip().replace(',', '')
    if not re.fullmatch(r'-?\d+', t):
        raise CellError('セル(%d,%d)が数字でない: %r' % (r, c, v))
    return int(t)


def read_year(path, warn):
    wb = openpyxl.load_workbook(path, data_only=True)
    if '4' not in wb.sheetnames:
        warn.append('%s: 第4表のシートが無い' % os.path.basename(path))
        return None
    ws = wb['4']
    # 見出し行(3行目)の結合セル -> 先頭列だけが「新受」
    groups = []
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= 3 <= rng.max_row:
            v = ws.cell(rng.min_row, rng.min_col).value
            if v:
                groups.append((rng.min_col, str(v).strip()))
    groups.sort()
    courts = [(c, n) for c, n in groups if n in COURT2PREF]
    zenkoku = [c for c, n in groups if n == '全国総数']
    if len(courts) != 50:
        warn.append('%s: 地裁の列が%d本しかない（50本のはず）' % (os.path.basename(path), len(courts)))
        return None
    if not zenkoku:
        warn.append('%s: 「全国総数」の列が見つからない（検算できない）' % os.path.basename(path))
        return None
    zcol = zenkoku[0]

    lab = {}
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        b = ws.cell(r, 3).value
        if a:
            lab[r] = (str(a).strip(), str(b).strip() if b else None)

    out = {}
    national = {}
    for label, sub, key in WANT:
        rs = [k for k, v in lab.items() if v[0] == label and (sub is None or v[1] == sub)]
        if not rs:
            warn.append('%s: 「%s」の行が無い' % (os.path.basename(path), label))
            return None
        r = rs[0]
        per = {}
        for c, n in courts:
            p = COURT2PREF[n]
            per[p] = per.get(p, 0) + cellint(ws, r, c)
        z = cellint(ws, r, zcol)
        # A. 50地裁の合計＝全国総数
        if sum(per.values()) != z:
            warn.append('%s: %s の地裁合計%d が全国総数%d と合わない（列ずれ）'
                        % (os.path.basename(path), label, sum(per.values()), z))
            return None
        out[key] = per
        national[key] = z
    return {'pref': out, 'national': national}


def read_shizen(path, warn, year, zenkoku_hasan):
    """第105表（破産新受事件数―受理区分別―全地方裁判所）から自然人・法人の内訳を読む。
    ★この表は**全国計しかない**（地裁別は第4表にしか無い）。だから「◯県の自己破産のうち
      何件が個人か」は公表されていない。記事では全国の比率だけを出し、
      県別の数字は『法人を含む』と断って出す。
    行は 総数／自然人／うち自己破産／法人・その他／うち自己破産／開始決定 の6つ。
    関所: 自然人＋法人＝総数＝第4表の全国総数。合わなければ採らない。"""
    wb = openpyxl.load_workbook(path, data_only=True)
    if '105' not in wb.sheetnames:
        warn.append('%d年: 第105表が無い（自然人の割合は出せない）' % year)
        return None
    ws = wb['105']
    for r in range(1, ws.max_row + 1):
        vals = []
        for c in range(1, min(ws.max_column, 12) + 1):
            try:
                vals.append(cellint(ws, r, c))
            except CellError:
                vals = []
                break
        vals = [v for v in vals if v]
        if len(vals) >= 6 and vals[0] > 1000:
            sou, shizen, shizen_jiko, hojin = vals[0], vals[1], vals[2], vals[3]
            if shizen + hojin != sou or sou != zenkoku_hasan:
                warn.append('%d年: 第105表の内訳が合わない（自然人%d＋法人%d≠総数%d／第4表%d）'
                            % (year, shizen, hojin, sou, zenkoku_hasan))
                return None
            return {'sou': sou, 'shizen': shizen, 'shizen_jiko': shizen_jiko, 'hojin': hojin}
    warn.append('%d年: 第105表から数字の行を見つけられない' % year)
    return None


def read_bar(path, warn):
    t = '\n'.join((p.extract_text() or '') for p in pypdf.PdfReader(path).pages)
    m = re.search(r'（\s*([０-９\d]{4})\s*年\s*([０-９\d]{1,2})\s*月\s*([０-９\d]{1,2})\s*日現在', t)
    z2h = lambda s: s.translate(str.maketrans('０１２３４５６７８９', '0123456789'))
    asof = '%s-%02d-%02d' % (z2h(m.group(1)), int(z2h(m.group(2))), int(z2h(m.group(3)))) if m else None
    if not asof:
        warn.append('弁護士会別会員数: 基準日が読めない')
    per = {}
    kai = 0
    for line in t.split('\n'):
        mm = re.match(r'^\s*(\S+)\s+(\d+)\s+(\d+)\s+[\d.]+%', line)
        if not mm or mm.group(1) not in BAR2PREF:
            continue
        per[BAR2PREF[mm.group(1)]] = per.get(BAR2PREF[mm.group(1)], 0) + int(mm.group(2))
        kai += 1
    mt = re.search(r'合計\s+(\d+)\s+(\d+)\s+[\d.]+%', t)
    total = int(mt.group(1)) if mt else None
    # C. 52会そろい、合計が一致すること
    if kai != 52:
        warn.append('弁護士会別会員数: %d会しか読めていない（52会のはず）' % kai)
    if total is not None and sum(per.values()) != total:
        warn.append('弁護士会別会員数: 合計が合わない（読み取り%d／PDF%d）' % (sum(per.values()), total))
    return {'asof': asof, 'total': total, 'per': per, 'kai': kai}


def corr(a, b):
    ma = sum(a) / len(a)
    mb = sum(b) / len(b)
    d = (sum((x - ma) ** 2 for x in a) * sum((y - mb) ** 2 for y in b)) ** .5
    return sum((x - ma) * (y - mb) for x, y in zip(a, b)) / d if d else None


def main():
    warn = []
    years = {}
    for path in sorted(glob.glob(os.path.join(CACHE, 'shiho-*.xlsx'))):
        y = int(re.search(r'shiho-(\d{4})', path).group(1))
        try:
            got = read_year(path, warn)
        except CellError as e:
            warn.append('%d年: %s → この年は採らない' % (y, e))
            got = None
        if got:
            got['shizen'] = read_shizen(path, warn, y, got['national']['hasan'])
            years[y] = got
            print('ok   %d年  破産%d 小規模個人再生%d' % (y, got['national']['hasan'], got['national']['saisei']))
    if not years:
        print('採れた年が無い', file=sys.stderr)
        for w in warn:
            print('  ! ' + w, file=sys.stderr)
        sys.exit(1)
    latest = max(years)

    bar = read_bar(os.path.join(CACHE, 'bengoshi.pdf'), warn)

    # 人口・被保護人員は生活保護の表から都道府県ごとに足す
    with io.open(HOGO, encoding='utf-8') as f:
        H = json.load(f)
    pop, hogo = {}, {}
    for r in H['rows']:
        p = r['name'] if r['level'] == '都道府県' else r['pref']
        if not p:
            continue
        pop[p] = pop.get(p, 0) + (r['jinko'] or 0)
        hogo[p] = hogo.get(p, 0) + (r['hogo_nin'] or 0)
    if len(pop) != 47:
        warn.append('人口を作れた都道府県が%d件しかない（47件のはず）' % len(pop))

    cur = years[latest]['pref']
    rows = []
    for p in sorted(pop):
        n_hasan = cur['hasan'].get(p, 0)
        n_saisei = cur['saisei'].get(p, 0) + cur['kyuyo'].get(p, 0)
        row = {
            'pref': p,
            'pop': pop[p],
            'bengoshi': bar['per'].get(p, 0),
            'hogo_nin': hogo.get(p, 0),
            'hasan': n_hasan,
            'saisei': cur['saisei'].get(p, 0),
            'kyuyo': cur['kyuyo'].get(p, 0),
            'hist': {},
        }
        for y in sorted(years):
            pr = years[y]['pref']
            row['hist'][str(y)] = {
                'hasan': pr['hasan'].get(p, 0),
                'saisei': pr['saisei'].get(p, 0) + pr['kyuyo'].get(p, 0),
            }
        rows.append(row)

    # 相関（記事の主張はここで出た数字だけを使う）
    PR = [r['pref'] for r in rows]
    hr = [r['hasan'] / r['pop'] * 100000 for r in rows]
    sr = [(r['saisei'] + r['kyuyo']) / r['pop'] * 100000 for r in rows]
    br = [r['bengoshi'] / r['pop'] * 100000 for r in rows]
    gr = [r['hogo_nin'] / r['pop'] * 1000 for r in rows]
    sh = [(r['saisei'] + r['kyuyo']) / (r['hasan'] + r['saisei'] + r['kyuyo']) * 100 for r in rows]
    c = {
        'hasan_hogo': corr(hr, gr),
        'hasan_bengoshi': corr(hr, br),
        'saisei_hogo': corr(sr, gr),
        'saisei_bengoshi': corr(sr, br),
        'share_hogo': corr(sh, gr),
        'share_bengoshi': corr(sh, br),
        'bengoshi_hogo': corr(br, gr),
    }
    # 偏相関: 生活保護率を固定したときの「破産率 と 弁護士密度」
    rhb, rhg, rbg = c['hasan_bengoshi'], c['hasan_hogo'], c['bengoshi_hogo']
    c['hasan_bengoshi_partial'] = (rhb - rhg * rbg) / (((1 - rhg ** 2) * (1 - rbg ** 2)) ** .5)

    data = {
        'source': {
            'shiho': '最高裁判所「司法統計年報 1.民事・行政編」第4表（民事・行政事件数―事件の種類及び新受、既済、未済―全地方裁判所及び地方裁判所別）',
            'shiho_url': 'https://www.courts.go.jp/toukei_siryou/shihotokei_nenpo/index.html',
            'bar': '日本弁護士連合会「弁護士会別会員数」',
            'bar_url': 'https://www.nichibenren.or.jp/jfba_info/membership/about.html',
            'pop': '人口・被保護人員は厚生労働省「被保護者調査」第9表からの逆算（0.5%ほどの丸め誤差を含む）',
            'note': '年は暦年（1〜12月）。年度ではない。数字は新受件数。破産には法人が含まれる。',
        },
        'latest_year': latest,
        'years': sorted(years),
        'bar_asof': bar['asof'],
        'bar_total': bar['total'],
        'national': {str(y): dict(years[y]['national'], shizen=years[y].get('shizen')) for y in sorted(years)},
        'corr': c,
        'rows': rows,
        'warnings': warn,
    }
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with io.open(OUT, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=1)
    print('書いた: %s  (%d年 47都道府県 / 弁護士は%s現在)' % (OUT, len(years), bar['asof']))
    for k, v in c.items():
        print('   相関 %-24s %+.3f' % (k, v))
    for w in warn:
        print('  ! ' + w)


if __name__ == '__main__':
    main()
